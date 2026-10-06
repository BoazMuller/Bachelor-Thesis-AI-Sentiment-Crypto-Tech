"""Map the supplied revised paper to the nine feedback points in feedback DG."""
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'output/docx/reviewer_changes_summary.docx'
ROWS = [
    ('1  Reviewer 1', 'Weekend and holiday sentiment',
     'Section 3 and Section 3.1\nPDF pp. 13-14\nAppendix A.1, PDF p. 48',
     'As proposed.', ' Explains rolling non-trading-day texts into the next trading day. Clarifies that individual texts are pooled before source averages are calculated; the appendix is aligned with this procedure.'),
    ('2  Reviewer 1', 'Prediction market measures',
     'Section 3.3\nPDF pp. 19-20',
     'Implemented with clarification.', ' Specifies the OpenAI AGI-announcement contract and defines the Metaculus transformation as negative days to the forecast date, not a reciprocal. Retains the distinction between trading and community forecasts.'),
    ('3  Reviewer 1', 'S&P 500 control choice',
     'Section 4.2\nPDF p. 23',
     'Modified explanation.', ' Explains the different purposes of the two regression stages. Omits the proposed near-collinearity rationale; VIX and equity returns are described as different controls. No extra regression is added.'),
    ('4  Reviewer 1', 'Short sample period',
     'Section 6\nPDF p. 43',
     'Implemented with an omission.', ' Adds the sample-length limitation and N = 487-501. Omits the proposed comparison with sample sizes in other studies and does not suggest that Kalman filtering guarantees statistical power.'),
    ('5  Reviewer 2', 'System size and Bitcoin spillovers',
     'Section 3.2, PDF p. 18\nSection 5.2, PDF pp. 26-28\nTable 4 Panel C, PDF p. 27\nTable 5, PDF p. 28\nSection 5.4.5, PDF pp. 37-39\nFigure 3, PDF p. 38; Table 9, PDF p. 39\nSection 6, PDF p. 40',
     'Implemented with qualifications.', ' Uses the proposed firms and corpus check, but identifies the basket as unmatched and acknowledges Cisco\'s AI exposure. The comparison has lower average TCI but a significant positive sentiment coefficient, qualifying AI exclusivity. The pairwise table omits the proposed claim of system invariance.'),
    ('6  Reviewer 2', 'EAIS and residualization',
     'Section 3.1\nPDF pp. 16-17',
     'As proposed.', ' Reframes EAIS as macro- and expectation-adjusted sentiment. Explains the insignificant prediction-market coefficients and avoids interpreting the residual itself as identified irrational beliefs.'),
    ('7  Reviewer 2', 'Classifier validation and technology versus investor sentiment',
     'Section 2.7, PDF p. 12\nSection 3.1, PDF p. 15',
     'Implemented with a qualification.', ' Retains the 100-text validation as an indicative check. Presents attention as a possible channel, not an identified mechanism. The transmission subsection is numbered 2.7 in the paper, rather than 2.6 in the proposal.'),
    ('8  Reviewer 2', 'Causality and same-day results',
     'Section 5.3, PDF pp. 30-34\nSection 6, PDF pp. 40-41\nTable A8, PDF p. 54',
     'Implemented with scope limits.', ' Adds the Granger tests and targeted causal-language edits in Results and Discussion. Stronger causal or speculative claims remain in the abstract and earlier sections. The new table is A8 rather than the proposed A7b.'),
    ('9  Reviewer 2', 'Raw sentiment robustness',
     'Section 5.4.1, PDF p. 35\nSection 5.4.3, PDF p. 36\nSection 6, PDF pp. 41-42',
     'Modified interpretation.', ' Treats AIS-EAIS divergence as sensitivity to filtering, not validation of a purified speculative component. Distinguishes the unstable BTC model from insignificant NASDAQ-100 estimates.'),
]

doc = Document()
# Remove the bundled template's inherited title rule as well as direct borders.
for root in [doc._element, doc.styles.element]:
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
sec = doc.sections[0]
sec.orientation = WD_ORIENT.LANDSCAPE
sec.page_width = Inches(11.69)
sec.page_height = Inches(8.27)
sec.top_margin = Inches(.55)
sec.bottom_margin = Inches(.55)
sec.left_margin = Inches(.60)
sec.right_margin = Inches(.60)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(4)
normal.paragraph_format.line_spacing = 1.04
title = doc.add_paragraph('Reviewer feedback implementation summary', 'Title')
title.paragraph_format.space_after = Pt(6)
for r in title.runs:
    r.font.name = 'Calibri';r.font.size = Pt(20);r.font.color.rgb = RGBColor(0,0,0)
doc.add_paragraph('This table maps the nine feedback points to the yellow highlights in the revised paper and records concise departures from the proposed response.')
p = doc.add_paragraph('Page references use the PDF viewer numbering, with the cover as page 1. Yellow marks added or revised passages and new tables or figures; unchanged relocated text, formatting and automatic renumbering are excluded.')
p.paragraph_format.space_after=Pt(8)
for r in p.runs:r.font.size=Pt(9)

table=doc.add_table(rows=1, cols=3)
table.alignment=WD_TABLE_ALIGNMENT.CENTER
table.autofit=False
widths=[2.05,3.05,5.39]
for cell,width in zip(table.rows[0].cells,widths):cell.width=Inches(width)
for col,width in zip(table.columns,widths):col.width=Inches(width)
for cell,text in zip(table.rows[0].cells,['Feedback point','Yellow highlighted locations','Implementation against the proposal']):
    cell.text=text
    shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'E4E8ED');cell._tc.get_or_add_tcPr().append(shd)
    for run in cell.paragraphs[0].runs:run.bold=True
hdr=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(hdr)

for index,(number,topic,location,status,detail) in enumerate(ROWS):
    cells=table.add_row().cells
    for cell,width in zip(cells,widths):cell.width=Inches(width)
    p=cells[0].paragraphs[0];p.add_run(number).bold=True;p.add_run('\n'+topic)
    p=cells[1].paragraphs[0]
    for i,line in enumerate(location.split('\n')):
        if i:p.add_run('\n')
        r=p.add_run(line);r.font.highlight_color=WD_COLOR_INDEX.YELLOW
    p=cells[2].paragraphs[0];p.add_run(status).bold=True;p.add_run(detail)
    # A row remains together; the header repeats on continuation pages.
    cant=OxmlElement('w:cantSplit');table.rows[-1]._tr.get_or_add_trPr().append(cant)
    if index%2:
        for cell in cells:
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),'F7F8FA');cell._tc.get_or_add_tcPr().append(shd)

for row in table.rows:
    for cell in row.cells:
        cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        tcpr=cell._tc.get_or_add_tcPr()
        margins=OxmlElement('w:tcMar')
        for side in ['top','bottom','left','right']:
            item=OxmlElement('w:'+side);item.set(qn('w:w'),'100');item.set(qn('w:type'),'dxa');margins.append(item)
        tcpr.append(margins)
        borders=OxmlElement('w:tcBorders')
        for side in ['top','bottom','left','right']:
            item=OxmlElement('w:'+side);item.set(qn('w:val'),'single');item.set(qn('w:sz'),'4');item.set(qn('w:color'),'D9D9D9');borders.append(item)
        tcpr.append(borders)
        for p in cell.paragraphs:
            p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.02

p=doc.add_paragraph('Other highlighted edits: Section 5.2 (PDF p. 26) contains small changes to the benchmark and spillover wording. These are separate from the nine feedback points.')
p.paragraph_format.space_before=Pt(7)
for r in p.runs:r.font.size=Pt(9)
doc.core_properties.title='Reviewer feedback implementation summary'
doc.core_properties.subject='Locations and implementation of the nine feedback points'
doc.core_properties.author='Boaz Muller'
doc.save(OUT)
print(OUT)
